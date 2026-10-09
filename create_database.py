import re
import shutil
from collections import Counter
from pathlib import Path
from pypdf import PdfReader
from docx import Document as DocxFile
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_chroma import Chroma

DATA_PATH="data/books"
CHROMA_PATH="chroma"


def load_text(path):
    # .md and .txt are plain text: read the whole file as one document.
    # Collapse runs of spaces (text copied from a terminal is often padded) so chunks hold real content.
    text = re.sub(r"[ \t]+", " ", path.read_text(encoding="utf-8"))
    return [Document(page_content=text, metadata={"source": str(path)})]


def load_pdf(path):
    # PDFs are binary: extract text page by page so each chunk remembers its page number
    pages = [(page.extract_text() or "").splitlines() for page in PdfReader(path).pages]

    # Headers/footers (book title, website) repeat on every page and add noise to search.
    # Drop any line that appears on more than half the pages, plus bare page numbers.
    line_counts = Counter(line.strip() for lines in pages for line in set(lines) if line.strip())
    repeated = {line for line, count in line_counts.items() if len(pages) > 2 and count > len(pages) / 2}

    documents = []
    for page_number, lines in enumerate(pages, start=1):
        kept = [line for line in lines if line.strip() not in repeated and line.strip() != str(page_number)]
        text = re.sub(r"[ \t]+", " ", "\n".join(kept)).strip()
        if text:  # skip blank or image-only pages
            documents.append(Document(page_content=text, metadata={"source": str(path), "page": page_number}))
    return documents


def load_docx(path):
    # Word files: join the text of every non-empty paragraph and table cell into one document
    file = DocxFile(path)
    parts = [p.text for p in file.paragraphs if p.text.strip()]
    for table in file.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    text = re.sub(r"[ \t]+", " ", "\n\n".join(parts))
    return [Document(page_content=text, metadata={"source": str(path)})]


# Which loader handles which file extension. To support a new type, add a line here.
LOADERS = {
    ".md": load_text,
    ".txt": load_text,
    ".pdf": load_pdf,
    ".docx": load_docx,
}

# 1. Load every supported file in data/books (including subfolders)
documents = []
for path in sorted(Path(DATA_PATH).rglob("*")):
    loader = LOADERS.get(path.suffix.lower())
    if loader is None:
        continue
    loaded = loader(path)
    print(f"Loaded {path.name}: {len(loaded)} document(s)")
    documents.extend(loaded)

# 2. Split into chunks
splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=200, add_start_index=True)
chunks = splitter.split_documents(documents)
print(f"Split {len(documents)} document(s) into {len(chunks)} chunks")

# 3 + 4. Embed and save (start fresh each run)
shutil.rmtree(CHROMA_PATH, ignore_errors=True)
embeddings = HuggingFaceEmbeddings(model_name="sentence-transformers/all-MiniLM-L6-v2")
Chroma.from_documents(chunks, embeddings, persist_directory=CHROMA_PATH)
print(f"Saved to {CHROMA_PATH}/")
